"""Zephyra Lite — File parsers."""

import io
from dataclasses import dataclass
from typing import Any

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

@dataclass
class ParsedSection:
    text: str
    metadata: dict[str, Any]

def parse_txt(content: bytes) -> list[ParsedSection]:
    """Parse a plain text file."""
    # Try different encodings
    encodings = ["utf-8", "windows-1252", "latin-1"]
    text = ""
    for enc in encodings:
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
            
    if not text:
        raise ValueError("Could not decode text file.")
        
    return [ParsedSection(text=text, metadata={})]

def parse_pdf(content: bytes) -> list[ParsedSection]:
    """Parse a PDF file page by page."""
    reader = PdfReader(io.BytesIO(content))
    sections = []
    
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            sections.append(ParsedSection(text=text.strip(), metadata={"page": i + 1}))
            
    if not sections:
        raise ValueError("Unable to extract readable text from this PDF.")
        
    return sections

def parse_docx(content: bytes) -> list[ParsedSection]:
    """Parse a DOCX file."""
    doc = Document(io.BytesIO(content))
    text_blocks = []
    
    for para in doc.paragraphs:
        txt = para.text.strip()
        if txt:
            text_blocks.append(txt)
            
    if not text_blocks:
        raise ValueError("Unable to extract readable text from this DOCX.")
        
    return [ParsedSection(text="\n\n".join(text_blocks), metadata={})]

def parse_xlsx(content: bytes) -> list[ParsedSection]:
    """Parse an XLSX workbook."""
    wb = load_workbook(filename=io.BytesIO(content), read_only=True, data_only=True)
    sections = []
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = []
        for row in sheet.iter_rows(values_only=True):
            # Convert non-None cells to string
            row_vals = [str(cell) if cell is not None else "" for cell in row]
            # Only add if row has some data
            if any(val.strip() for val in row_vals):
                rows.append(" | ".join(row_vals))
                
        if rows:
            text = f"Sheet: {sheet_name}\n" + "\n".join(rows)
            sections.append(ParsedSection(text=text, metadata={"sheet": sheet_name}))
            
    if not sections:
        raise ValueError("Unable to extract readable data from this XLSX.")
        
    return sections

def parse_document(content: bytes, mime_type: str, filename: str) -> list[ParsedSection]:
    """Parse a document based on its MIME type or extension."""
    if mime_type == "application/pdf" or filename.lower().endswith(".pdf"):
        return parse_pdf(content)
    elif mime_type in ["text/plain", "text/csv"] or filename.lower().endswith((".txt", ".csv", ".md")):
        return parse_txt(content)
    elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or filename.lower().endswith(".docx"):
        return parse_docx(content)
    elif mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or filename.lower().endswith(".xlsx"):
        return parse_xlsx(content)
    else:
        raise ValueError("Unsupported file type. Supported formats are PDF, TXT, DOCX, and XLSX.")
